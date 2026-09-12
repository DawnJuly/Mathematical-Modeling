# -*- coding: utf-8 -*-
"""Q3/Q4 公共核心：数据加载、时段LP、光伏预报展开(约定B)、MPC滚动、费用模型A、写xlsx
约定B（已由 explore_fc2/explore_fc3 全年统计验证）：附件3「预报k小时」对应
发布时刻起的第 k 个小时区间 [T+k-1, T+k)，即 0:00 发布的 预报1..24小时 恰好
覆盖当天 0:00..24:00 共 24 个小时；小时值(kW)阶梯应用到该小时 6 个 10 分钟时段。"""
import numpy as np
import pandas as pd
from pathlib import Path
from scipy.optimize import linprog
from scipy.sparse import lil_matrix

BASE = Path(__file__).resolve().parent.parent

DT = 1.0 / 6.0
T = 144
PMAX = 5000.0 * DT          # 833.3333 kWh / 10min
ETA = 0.9
SOC_LO, SOC_HI = 1200.0, 10800.0
EMERG = 5.0
UP_FEE, DOWN_FEE = 1.5, 0.5
STAGES = [0, 6, 12, 18]     # 发布/调整时刻(小时)
OUT_START = 31              # 2025-02-01 的日索引
DAYS_OUT = list(range(OUT_START, 365))

BLOCKS = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
BNAMES = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00',
          '16:00-20:00', '20:00-24:00']
SIX_TIMES = [('10:00', 60), ('12:00', 72), ('14:00', 84),
             ('16:00', 96), ('18:00', 108), ('20:00', 120)]
FOUR_DAYS = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']


def load_data():
    a1 = pd.read_excel(BASE / '附件/附件1.xlsx')
    price1 = a1['电价'].to_numpy(float)                       # (144,)
    load_df = pd.read_excel(BASE / '附件/附件2.xlsx', sheet_name='小区负载')
    pv_df = pd.read_excel(BASE / '附件/附件2.xlsx', sheet_name='光伏发电实际功率')
    dates = pd.to_datetime(load_df.iloc[:, 0])
    LOAD = load_df.iloc[:, 1:].to_numpy(float) * DT           # (365,144) kWh
    PV = pv_df.iloc[:, 1:].to_numpy(float) * DT
    a4 = pd.read_excel(BASE / '附件/附件4.xlsx')
    PRICE4 = a4.iloc[:, 1:].to_numpy(float)                   # (365,144)
    a3 = pd.read_excel(BASE / '附件/附件3.xlsx')
    a3.iloc[:, 0] = a3.iloc[:, 0].ffill()
    d3 = pd.to_datetime(a3.iloc[:, 0])
    # FC[d, j, k-1]: 第 d 天 STAGES[j] 时刻发布的「预报k小时」, 单位 kW(小时均值)
    FC = np.zeros((365, 4, 24))
    didx = {pd.Timestamp(x): i for i, x in enumerate(pd.unique(d3))}
    for j, h0 in enumerate(STAGES):
        sub = a3[a3.iloc[:, 1] == '%d:00' % h0]
        assert len(sub) == 365
        for _, row in sub.iterrows():
            FC[didx[pd.Timestamp(row.iloc[0])], j] = row.iloc[2:].to_numpy(float)
    n_neg = int((FC < 0).sum())
    if n_neg:
        FC = np.maximum(FC, 0.0)
    HA = PV.reshape(365, 24, 6).mean(axis=2) / DT             # 实际逐时功率 kW
    return dict(price1=price1, dates=dates, LOAD=LOAD, PV=PV,
                PRICE4=PRICE4, FC=FC, HA=HA, n_neg_fc=n_neg)


def fc_slots(FC, d, j):
    """第 d 天 STAGES[j] 时刻发布的预报 -> 当天 144 时段光伏能量(kWh)。
    仅覆盖 h0=STAGES[j] 之后的小时 (h0..23); h0 之前的时段为 0(不会用到)。"""
    h0 = STAGES[j]
    f = np.zeros(24)
    for k in range(1, 25 - h0):
        f[h0 + k - 1] = FC[d, j, k - 1]
    return np.repeat(f, 6) * DT


# ---------- 时段 LP: 变量 buy[n], ch[n], dis[n], e[n], soc[n+1] ----------
_cache = {}

def solve_lp(price, pv_e, load_e, soc0):
    """从某时段起到当天结束(n=len(price)个时段)的购电优化。
    price/pv_e/load_e 为剩余时段数组。返回 buy,ch,dis,e,soc。"""
    n = len(price)
    if n not in _cache:
        N = 5 * n + 1
        Aeq = lil_matrix((n, N))
        for t in range(n):
            Aeq[t, 4 * n + t + 1] = 1.0
            Aeq[t, 4 * n + t] = -1.0
            Aeq[t, n + t] = -ETA
            Aeq[t, 2 * n + t] = 1.0 / ETA
        Aub = lil_matrix((n, N))
        for t in range(n):
            Aub[t, t] = -1.0
            Aub[t, 2 * n + t] = -1.0
            Aub[t, 3 * n + t] = -1.0
            Aub[t, n + t] = 1.0
        _cache[n] = (Aeq.tocsr(), Aub.tocsr())
    Aeq, Aub = _cache[n]
    c = np.concatenate([price, np.zeros(2 * n), EMERG * price, np.zeros(n + 1)])
    bounds = ([(0, None)] * n + [(0, PMAX)] * (2 * n) + [(0, None)] * n +
              [(SOC_LO, SOC_HI)] * (n + 1))
    bounds[4 * n] = (soc0, soc0)
    res = linprog(c, A_ub=Aub, b_ub=pv_e - load_e, A_eq=Aeq,
                  b_eq=np.zeros(n), bounds=bounds, method='highs')
    assert res.status == 0, res.message
    x = res.x
    return (x[0:n], x[n:2 * n], x[2 * n:3 * n], x[3 * n:4 * n], x[4 * n:5 * n + 1])


def q2_year(price_mat):
    """理想情形(光伏已知): 逐日 LP, SOC 跨天传递。返回 dict(day->dict)。"""
    D = load_data()
    PV, LOAD = D['PV'], D['LOAD']
    out = {}
    soc0 = 6000.0
    for d in range(365):
        pr = price_mat[d] if price_mat.ndim == 2 else price_mat
        buy, ch, dis, e, soc = solve_lp(pr, PV[d], LOAD[d], soc0)
        out[d] = dict(buy=buy, ch=ch, dis=dis, e=e, soc=soc,
                      cost=float(pr @ buy + EMERG * pr @ e))
        soc0 = float(soc[-1])
    return out


def q3_year(price_mat, FC, LOAD, PV, adjust=True):
    """MPC 滚动: 0:00 计划 + 6/12/18:00 三次调整(adjust=False 则为对照组)。
    返回 dict(day->dict), 含 plan/adj/ch/dis/e/soc 及费用四项分解。"""
    out = {}
    soc0 = 6000.0
    for d in range(365):
        pr = price_mat[d] if price_mat.ndim == 2 else price_mat
        pv_fc0 = fc_slots(FC, d, 0)
        plan_buy, plan_ch, plan_dis, _, plan_soc = solve_lp(pr, pv_fc0, LOAD[d], soc0)
        buy = plan_buy.copy(); ch = plan_ch.copy(); dis = plan_dis.copy()
        if adjust:
            s_now = soc0
            for j, h0 in enumerate(STAGES[1:], start=1):
                t0 = h0 * 6
                s_now = soc0 + ETA * ch[:t0].sum() - dis[:t0].sum() / ETA
                pv_f = fc_slots(FC, d, j)
                b2, c2, d2, _, _ = solve_lp(pr[t0:], pv_f[t0:], LOAD[d, t0:], s_now)
                buy[t0:] = b2; ch[t0:] = c2; dis[t0:] = d2
        # 实际交割: 缺口 = 紧急购电
        resid = LOAD[d] + ch - buy - dis - PV[d]
        e = np.maximum(resid, 0.0)
        e[e < 1e-9] = 0.0
        soc = np.empty(T + 1)
        soc[0] = soc0
        soc[1:] = soc0 + ETA * np.cumsum(ch) - np.cumsum(dis) / ETA
        plan_cost = float(pr @ plan_buy)
        up = float((UP_FEE * pr * np.maximum(buy - plan_buy, 0)).sum())
        down = float((DOWN_FEE * pr * np.maximum(plan_buy - buy, 0)).sum())
        emerg = float((EMERG * pr * e).sum())
        out[d] = dict(plan=plan_buy, plan_ch=plan_ch, plan_dis=plan_dis,
                      adj=buy, ch=ch, dis=dis, e=e, soc=soc,
                      plan_cost=plan_cost, up=up, down=down, emerg=emerg,
                      total=plan_cost + up - down + emerg)
        soc0 = float(soc[-1])
    return out


def emerg_intervals(e):
    """把 e>1e-6 的连续时段合并为 (起始slot, 结束slot不含, 电量) 列表"""
    idx = np.where(e > 1e-6)[0]
    if len(idx) == 0:
        return []
    segs, s, prev = [], idx[0], idx[0]
    for t in idx[1:]:
        if t == prev + 1:
            prev = t
        else:
            segs.append((s, prev + 1, float(e[s:prev + 1].sum())))
            s = prev = t
    segs.append((s, prev + 1, float(e[s:prev + 1].sum())))
    return segs


def slot_lab(t):
    h, m = t // 6, (t % 6) * 10
    return '%d:%02d' % (h, m)


def write_result(path_tpl, path_out, dates, res, with_adj, extra_sheet_date_col=True):
    """复制模板并填写。res: dict(day->dict)。返回写日志用信息。"""
    import shutil
    from openpyxl import load_workbook
    shutil.copy(path_tpl, path_out)
    wb = load_workbook(path_out)
    n_emerg = 0
    # 计划购电量
    ws = wb['计划购电量']
    for i, d in enumerate(DAYS_OUT):
        r, row = res[d], i + 2
        for t in range(T):
            ws.cell(row=row, column=2 + t, value=round(float(r['plan'][t]), 4))
        ws.cell(row=row, column=146, value=round(float(r['plan'].sum()), 4))
        ws.cell(row=row, column=147, value=round(float(r['plan_cost']), 4))
    # 调整购电量
    if with_adj:
        ws = wb['调整购电量']
        for i, d in enumerate(DAYS_OUT):
            r, row = res[d], i + 2
            for t in range(T):
                ws.cell(row=row, column=2 + t, value=round(float(r['adj'][t]), 4))
            ws.cell(row=row, column=146, value=round(float(r['adj'].sum()), 4))
            # 全天购电费 = 当天实际总电费(计划+调整费-回退费+紧急购电)
            ws.cell(row=row, column=147, value=round(float(r['total']), 4))
    # 充放电量
    ws = wb['充放电量']
    ws.delete_rows(2, ws.max_row)
    row = 2
    for d in DAYS_OUT:
        r = res[d]
        for k, (a, b) in enumerate(BLOCKS):
            if k == 0:
                ws.cell(row=row, column=1, value=dates[d].to_pydatetime())
            ws.cell(row=row, column=2, value=BNAMES[k])
            ws.cell(row=row, column=3, value=round(float(r['ch'][a:b].sum()), 4))
            ws.cell(row=row, column=4, value=round(float(r['dis'][a:b].sum()), 4))
            if k == 0:
                ws.cell(row=row, column=5, value='0:00')
                ws.cell(row=row, column=6, value=round(float(r['soc'][0]), 4))
            elif k == 1:
                ws.cell(row=row, column=5, value='24:00')
                ws.cell(row=row, column=6, value=round(float(r['soc'][-1]), 4))
            row += 1
    # 紧急购电量: 连续非零时段合并
    ws = wb['紧急购电量']
    ws.delete_rows(2, ws.max_row)
    row = 2
    for d in DAYS_OUT:
        segs = emerg_intervals(res[d]['e'])
        first = True
        for s, t2, v in segs:
            if first:
                ws.cell(row=row, column=1, value=dates[d].to_pydatetime())
                first = False
            ws.cell(row=row, column=2, value='%s-%s' % (slot_lab(s), slot_lab(t2)))
            ws.cell(row=row, column=3, value=round(v, 4))
            row += 1
            n_emerg += 1
    wb.save(path_out)
    return n_emerg


def fc_error_stats(D):
    """附件3 全部预报 vs 附件2 实际(约定B): 按提前时长 k=1..24 统计。"""
    FC, HA = D['FC'], D['HA']
    per_k = {k: [] for k in range(1, 25)}
    per_k_day = {k: [] for k in range(1, 25)}
    allerr = []
    for d in range(365):
        for j, h0 in enumerate(STAGES):
            for k in range(1, 25):
                h_abs = h0 + k - 1
                dd = d + h_abs // 24
                if dd > 364:
                    continue
                h = h_abs % 24
                err = FC[d, j, k - 1] - HA[dd, h]
                per_k[k].append(err)
                allerr.append(err)
                if max(FC[d, j, k - 1], HA[dd, h]) > 100:
                    per_k_day[k].append(err)
    allerr = np.array(allerr)
    tbl = []
    for k in range(1, 25):
        e = np.array(per_k[k])
        ed = np.array(per_k_day[k]) if per_k_day[k] else np.array([np.nan])
        tbl.append((k, len(e), np.abs(e).mean(), np.sqrt((e ** 2).mean()),
                    np.abs(ed).mean() if len(ed) else np.nan))
    overall = (np.abs(allerr).mean(), np.sqrt((allerr ** 2).mean()), allerr.mean())
    return tbl, overall
