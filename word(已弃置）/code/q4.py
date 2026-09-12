# -*- coding: utf-8 -*-
"""Q4: 波动电价(附件4)下重算 Q2/Q3。
输出: result4-2.xlsx, result4-3.xlsx (并覆盖 附件/附件5/), q4_log.txt"""
import shutil
import time
import numpy as np
from pathlib import Path
from openpyxl import load_workbook
from lp_core import (load_data, q2_year, q3_year, write_result,
                     BLOCKS, BNAMES, SIX_TIMES, FOUR_DAYS, DAYS_OUT,
                     SOC_LO, SOC_HI)

BASE = Path(__file__).resolve().parent.parent

t_start = time.time()
L = []

def log(s=''):
    L.append(s)
    print(s)

BAK = BASE / '附件备份/附件5'  # 空提交模板(原始文件)
assert (BAK / 'result4-2.xlsx').exists() and (BAK / 'result4-3.xlsx').exists(), \
    '缺少空提交模板: %s' % BAK

D = load_data()
price1, dates = D['price1'], D['dates']
LOAD, PV, FC = D['LOAD'], D['PV'], D['FC']
PRICE4 = D['PRICE4']
log('附件4 波动电价: min=%.4f max=%.4f mean=%.4f 元/kWh' %
    (PRICE4.min(), PRICE4.max(), PRICE4.mean()))
log('附件1 固定电价: min=%.4f max=%.4f mean=%.4f 元/kWh' %
    (price1.min(), price1.max(), price1.mean()))

# ---------- Q2/Q3 固定电价基准(取自q3, 重算以对照) ----------
q2_f = q2_year(price1)
q2f_tot = sum(q2_f[d]['cost'] for d in DAYS_OUT)
q3_f = q3_year(price1, FC, LOAD, PV, adjust=True)
q3f_tot = sum(q3_f[d]['total'] for d in DAYS_OUT)
assert abs(q2f_tot - 12254765.72) < 1.0

# ---------- Q4: 波动电价 ----------
t0 = time.time()
q2_v_raw = q2_year(PRICE4)
log('Q4-Q2(波动) 求解耗时 %.1fs' % (time.time() - t0))
q2_v = {d: dict(r, plan=r['buy'], plan_cost=r['cost'])
        for d, r in q2_v_raw.items()}
t0 = time.time()
q3_v = q3_year(PRICE4, FC, LOAD, PV, adjust=True)
ctrl_v = q3_year(PRICE4, FC, LOAD, PV, adjust=False)
log('Q4-Q3(波动)+对照 求解耗时 %.1fs' % (time.time() - t0))

def year_sum(res, key_total='total'):
    c = dict(plan=0., up=0., down=0., emerg=0., total=0., e=0., nday=0)
    for d in DAYS_OUT:
        r = res[d]
        c['plan'] += r.get('plan_cost', r.get('cost', 0.))
        c['up'] += r.get('up', 0.); c['down'] += r.get('down', 0.)
        c['emerg'] += r.get('emerg', 0.)
        c['total'] += r.get(key_total, r.get('cost', 0.))
        c['e'] += r['e'].sum()
        c['nday'] += int(r['e'].sum() > 1e-6)
    return c

y2v = year_sum(q2_v)
y3v = year_sum(q3_v)
ycv = year_sum(ctrl_v)

# ---------- 四天明细 ----------
four_idx = [i for i, t in enumerate(dates)
            if t.strftime('%Y-%m-%d') in FOUR_DAYS]
for d in four_idx:
    key = dates[d].strftime('%Y-%m-%d')
    log('')
    log('===== %s (day idx %d) =====' % (key, d))
    r2 = q2_v[d]
    log('[Q2重算-波动电价] 全天购电量 %.4f kWh  全天购电费 %.4f 元  紧急 %.6f kWh' %
        (r2['plan'].sum(), r2['plan_cost'], r2['e'].sum()))
    for name, t in SIX_TIMES:
        log('  %s-%d:00 1小时购电: %.4f kWh' % (name, int(name[:2]) + 1,
                                              r2['plan'][t:t + 6].sum()))
    for k, (a, b) in enumerate(BLOCKS):
        log('  %-11s 充电 %.4f  放电 %.4f' %
            (BNAMES[k], r2['ch'][a:b].sum(), r2['dis'][a:b].sum()))
    log('  0:00 储电量 %.4f   24:00 储电量 %.4f   SOC min/max %.2f/%.2f' %
        (r2['soc'][0], r2['soc'][-1], r2['soc'].min(), r2['soc'].max()))
    r3 = q3_v[d]
    log('[Q3重算-波动电价] 计划 %.4f kWh/%.2f元 -> 调整后 %.4f kWh, 当天实际总电费 %.2f 元' %
        (r3['plan'].sum(), r3['plan_cost'], r3['adj'].sum(), r3['total']))
    for name, t in SIX_TIMES:
        log('  %s-%d:00 1小时: 计划 %.4f  调整后 %.4f kWh' %
            (name, int(name[:2]) + 1,
             r3['plan'][t:t + 6].sum(), r3['adj'][t:t + 6].sum()))
    for k, (a, b) in enumerate(BLOCKS):
        log('  %-11s 充电 %.4f  放电 %.4f' %
            (BNAMES[k], r3['ch'][a:b].sum(), r3['dis'][a:b].sum()))
    log('  0:00 储电量 %.4f   24:00 储电量 %.4f   SOC min/max %.2f/%.2f' %
        (r3['soc'][0], r3['soc'][-1], r3['soc'].min(), r3['soc'].max()))
    log('  费用分解: 计划 %.2f + 上调 %.2f - 下调回退 %.2f + 紧急 %.2f = %.2f 元' %
        (r3['plan_cost'], r3['up'], r3['down'], r3['emerg'], r3['total']))
    log('  紧急购电量: %.6f kWh   对照组(波动)当天总电费: %.2f 元' %
        (r3['e'].sum(), ctrl_v[d]['total']))

# ---------- 全年汇总 ----------
log('')
log('===== 全年汇总 (2025-02-01 ~ 2025-12-31, 334天) =====')
log('Q2重算(波动): 总电费 %.2f 元, 紧急购电 %.4f kWh (必须≈0)' %
    (y2v['total'], y2v['e']))
log('Q3重算(波动): 计划 %.2f + 上调 %.2f - 下调回退 %.2f + 紧急 %.2f = 总电费 %.2f 元' %
    (y3v['plan'], y3v['up'], y3v['down'], y3v['emerg'], y3v['total']))
log('  紧急购电总量 %.4f kWh, 非零天数 %d' % (y3v['e'], y3v['nday']))
log('对照组(波动): 计划 %.2f + 紧急 %.2f = 总电费 %.2f 元' %
    (ycv['plan'], ycv['emerg'], ycv['total']))
log('  紧急购电总量 %.4f kWh, 非零天数 %d' % (ycv['e'], ycv['nday']))
sv = ycv['total'] - y3v['total']
log('波动电价下滚动调整相对对照组节省: %.2f 元 (%.2f%%)' %
    (sv, 100 * sv / ycv['total']))
log('')
log('===== 固定电价 vs 波动电价 =====')
log('Q2: 固定 %.2f  波动 %.2f  差值 %.2f 元 (%.2f%%)' %
    (q2f_tot, y2v['total'], y2v['total'] - q2f_tot,
     100 * (y2v['total'] - q2f_tot) / q2f_tot))
log('Q3: 固定 %.2f  波动 %.2f  差值 %.2f 元 (%.2f%%)' %
    (q3f_tot, y3v['total'], y3v['total'] - q3f_tot,
     100 * (y3v['total'] - q3f_tot) / q3f_tot))

# ---------- 自验 ----------
log('')
log('===== 自验 =====')
assert y2v['e'] < 1e-6, 'Q4-Q2 紧急购电必须为0'
log('Q4-Q2 全年紧急购电 = %.2e kWh (≈0): 通过' % y2v['e'])
for tag, res, pr in [('Q4-Q2', q2_v, PRICE4), ('Q4-Q3', q3_v, PRICE4),
                     ('Q4-对照', ctrl_v, PRICE4)]:
    for d in range(365):
        r = res[d]
        assert r['soc'].min() >= SOC_LO - 1e-6, (tag, d)
        assert r['soc'].max() <= SOC_HI + 1e-6, (tag, d)
        buy = r.get('adj', r['buy'] if 'buy' in r else r['plan'])
        resid = buy + PV[d] + r['dis'] + r['e'] - LOAD[d] - r['ch']
        assert resid.min() > -1e-5, (tag, d, resid.min())
    log('%s: 365天 SOC界、能量平衡、e≥0 全部通过' % tag)
for d in range(365):
    pr = PRICE4[d]
    assert q3_v[d]['total'] >= pr @ q3_v[d]['adj'] - 1e-4
log('Q4-Q3 total>=Σprice·adj 下界: 通过')

# ---------- 写 xlsx ----------
OUT42 = BASE / 'result4-2.xlsx'
OUT43 = BASE / 'result4-3.xlsx'
n_em2 = write_result(BAK / 'result4-2.xlsx', OUT42,
                     dates, q2_v, with_adj=False)
shutil.copy(OUT42, BASE / '附件/附件5/result4-2.xlsx')
n_em3 = write_result(BAK / 'result4-3.xlsx', OUT43,
                     dates, q3_v, with_adj=True)
shutil.copy(OUT43, BASE / '附件/附件5/result4-3.xlsx')
log('result4-2.xlsx (紧急区间行 %d), result4-3.xlsx (紧急区间行 %d) 已生成并覆盖 附件/附件5/'
    % (n_em2, n_em3))

wb = load_workbook(OUT42, read_only=True)
for name in ['计划购电量', '充放电量', '紧急购电量']:
    ws = wb[name]
    log('result4-2 sheet %-6s: %d 行 x %d 列' % (name, ws.max_row, ws.max_column))
assert wb['计划购电量'].max_row == len(DAYS_OUT) + 1
assert wb['充放电量'].max_row == len(DAYS_OUT) * 6 + 1
wb.close()
wb = load_workbook(OUT43, read_only=True)
for name in ['计划购电量', '调整购电量', '充放电量', '紧急购电量']:
    ws = wb[name]
    log('result4-3 sheet %-6s: %d 行 x %d 列' % (name, ws.max_row, ws.max_column))
assert wb['计划购电量'].max_row == len(DAYS_OUT) + 1
assert wb['调整购电量'].max_row == len(DAYS_OUT) + 1
assert wb['充放电量'].max_row == len(DAYS_OUT) * 6 + 1
wb.close()
log('sheet结构校验通过')

log('')
log('总耗时 %.1fs' % (time.time() - t_start))
with open(BASE / 'q4_log.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(L))
print('q4_log.txt written')
