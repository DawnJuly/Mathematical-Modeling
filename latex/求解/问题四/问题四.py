# 问题四：波动电价下重算问题2、问题3
# 4-2 与问题二同口径：每天 0:00 预测 — 决策 — 结算 三段式
# 4-3 为滚动时域调整（日内 4 次预报更新）
import sys, os
HERE = os.path.dirname(os.path.abspath(__file__))
SOLVE = os.path.dirname(HERE)
sys.path.insert(0, SOLVE)
from _common import *
from _models import *

price_fixed, _, _ = read_附件1()
dates, load_all, pv_actual_all = read_附件2()
d4, price_all = read_附件4()
fc = read_附件3()

def date_key(d):
    return f'{d.year}-{d.month}-{d.day}'

def merge_runs(delta, thr=1e-6):
    """把非零紧急购电的连续间隔合并为 (起时刻, 止时刻, 电量kWh)。"""
    labs = interval_labels(); out = []; i = 0; n = len(delta)
    while i < n:
        if delta[i] > thr:
            j = i
            while j + 1 < n and delta[j + 1] > thr:
                j += 1
            out.append((labs[i].split('-')[0], labs[j].split('-')[1].replace('+1', ''),
                        float(np.sum(delta[i:j + 1]) * DT)))
            i = j + 1
        else:
            i += 1
    return out

N = len(dates)
S = 31   # 2025-02-01 索引；附件5仅要求保存 2.1—12.31，1月作为预热期，不纳入汇总统计
# ---- 4-2: 三段式（预测—随机规划—结算，波动电价） ----
g42 = np.zeros((N, N_INT)); ch42 = np.zeros((N, N_INT)); dch42 = np.zeros((N, N_INT))
dl42 = np.zeros((N, N_INT)); su42 = np.zeros((N, N_INT))
E42 = np.zeros((N, N_INT + 1))
plan42 = np.zeros(N); emg42 = np.zeros(N)
Lf42 = np.full((N, N_INT), np.nan); Sf42 = np.full((N, N_INT), np.nan)
res_pool_L, res_pool_P = [], []
PI = np.outer([0.2, 0.6, 0.2], [0.2, 0.6, 0.2]).ravel()   # 9 情景 = 3 负载 × 3 光伏
carry2 = E0_INIT
for d in range(N):
    if d == 0:
        # 1/1 0:00 尚无历史数据，计划量置零，储电量全天保持 6000 kWh
        E42[d] = np.full(N_INT + 1, E0_INIT)
        continue
    L_fc = forecast_load(load_all[:d], dates[:d], dates[d])
    S_fc = forecast_pv(pv_actual_all[:d])
    fL, fP = scenario_factors(np.concatenate(res_pool_L) if res_pool_L else np.array([]),
                              np.concatenate(res_pool_P) if res_pool_P else np.array([]))
    LF = np.repeat(fL, 3); SF = np.tile(fP, 3)
    g, ch, dch, E, _ = solve_daily_stochastic(L_fc, S_fc, price_all[d], LF, SF, PI, e0=carry2)
    delta, _, plan_cost, emerg_cost = settle_emergency(load_all[d], pv_actual_all[d], price_all[d], g, ch, dch)
    g42[d], ch42[d], dch42[d], E42[d] = g, ch, dch, E
    dl42[d] = delta; plan42[d] = plan_cost; emg42[d] = emerg_cost
    su42[d] = np.maximum(0.0, pv_actual_all[d] + g + dch - ch - load_all[d]) * DT
    Lf42[d], Sf42[d] = L_fc, S_fc
    carry2 = float(E[-1])
    res_pool_L.append(pred_residual(L_fc, load_all[d]))
    m = S_fc > 50.0
    if m.any():
        res_pool_P.append(((pv_actual_all[d] - S_fc) / np.maximum(S_fc, 1.0))[m])
cost42 = plan42 + emg42

gE42 = g42 * DT; chE42 = ch42 * DT; dchE42 = dch42 * DT
W = slice(S, N); nW = N - S
plan_kwh = float(gE42[W].sum()); plan_fee = float(plan42[W].sum())
emg_kwh = float(dl42[W].sum() * DT); emg_fee = float(emg42[W].sum())
tot_fee = plan_fee + emg_fee
n_emg_int = int((dl42[W].reshape(-1) > 1e-6).sum())
n_emg_day = int((dl42[W].sum(axis=1) > 1e-6).sum())
net_kwh = float((load_all[W] - pv_actual_all[W]).sum() * DT)
print(f'4-2 波动电价全年(365天)购电费={cost42.sum():.2f} 元')
print(f'[计划段] 计划购电量={plan_kwh:.2f} kWh, 计划购电费={plan_fee:.2f} 元')
print(f'[紧急段] 紧急购电量={emg_kwh:.2f} kWh, 紧急购电费={emg_fee:.2f} 元')
print(f'[合计]   4-2 2.1—12.31 总购电费={tot_fee:.2f} 元, 日均={tot_fee / nW:.2f} 元 '
      f'(固定电价问题2为 14382892.99 元)')
print(f'[紧急覆盖面] 非零间隔={n_emg_int} 个 ({n_emg_int / (nW * N_INT) * 100:.2f}%), '
      f'发生天数={n_emg_day} 天 ({n_emg_day / nW * 100:.2f}%)')
print(f'[紧急强度] 最大功率={dl42[W].max():.2f} kW, 单日最大电量={float((dl42[W].sum(axis=1) * DT).max()):.2f} kWh')
print(f'[储能] 2/1 0:00 SOC={E42[S, 0]:.2f}, 12/31 24:00 SOC={E42[N - 1, -1]:.2f}, '
      f'SOC 最小={E42[W, :].min():.2f}, SOC 最大={E42[W, :].max():.2f}')
print(f'[能量流] 净负荷={net_kwh:.2f} kWh, 放电量={float(dchE42[W].sum()):.2f} kWh, '
      f'充电量={float(chE42[W].sum()):.2f} kWh, 弃电={float(su42[W].sum()):.2f} kWh')
bal = net_kwh - plan_kwh - float(dchE42[W].sum()) + float(chE42[W].sum())
print(f'[自洽校验] 净负荷-计划-放电+充电={bal:.2f} = 紧急电量({emg_kwh:.2f}) - 弃电({float(su42[W].sum()):.2f})')

# ---- 4-3: 滚动（波动电价） ----
def run_day43(d):
    L = load_all[d]; S_act = pv_actual_all[d]; key = date_key(dates[d]); f = fc[key]
    pr = price_all[d]
    S0 = forecast_10min(0, f[0], 0, 143); S6 = forecast_10min(6, f[6], 0, 143)
    S12 = forecast_10min(12, f[12], 0, 143); S18 = forecast_10min(18, f[18], 0, 143)
    g0, ch0, dch0, E0, c_plan = solve_daily(L, S0, pr, e0=carry3[d])
    g6, ch6, dch6, E6, _ = solve_adjust(L, S6, pr, g0, e0=float(E0[36]), k0=36)
    g12, ch12, dch12, E12, _ = solve_adjust(L, S12, pr, g0, e0=float(E6[36]), k0=72)
    g18, ch18, dch18, E18, _ = solve_adjust(L, S18, pr, g0, e0=float(E12[36]), k0=108)
    blocks = [0, 36, 72, 108, 144]
    sols = [(g0, ch0, dch0, E0), (g6, ch6, dch6, E6), (g12, ch12, dch12, E12), (g18, ch18, dch18, E18)]
    adj = np.zeros(144); chf = np.zeros(144); dchf = np.zeros(144); Ef = np.empty(145); Ef[0] = E0[0]
    for i in range(4):
        a, b = blocks[i], blocks[i + 1]
        g_, c_, d_, E_ = sols[i]
        adj[a:b] = g_[0:b - a]; chf[a:b] = c_[0:b - a]; dchf[a:b] = d_[0:b - a]
        Ef[a + 1:b + 1] = E_[1:b - a + 1]
    plan = g0
    base = np.minimum(plan, adj) * DT * pr
    down = np.maximum(plan - adj, 0) * DT * pr * 0.5
    up = np.maximum(adj - plan, 0) * DT * pr * 1.5
    plan_commit = (plan * DT * pr).sum()
    purchase_total = (base + down + up).sum()
    # 紧急购电 = 实际供电缺口：delta_k = max{L - (g_k + S^r_k + d_k - c_k), 0}
    gap = np.maximum(L - (adj + S_act + dchf - chf), 0.0)
    emer_kwh = gap * DT
    emer_cost = (emer_kwh * 5.0 * pr).sum()
    total = purchase_total + emer_cost
    return dict(g0=g0, adj=adj, chf=chf, dchf=dchf, Ef=Ef, emer_kwh=emer_kwh, emer_cost=emer_cost,
                plan_commit=plan_commit, purchase_total=purchase_total, total=total, pr=pr)

carry3 = [E0_INIT] * N
rec43 = [None] * N
cost43 = np.zeros(N)
for d in range(N):
    r = run_day43(d)
    rec43[d] = r; cost43[d] = r['total']
    if d + 1 < N:
        carry3[d + 1] = float(r['Ef'][-1])
print(f'4-3 波动电价滚动全年(365天)总费用={cost43.sum():.2f} 元')
print(f'[汇总口径] 4-3 2.1—12.31 总费用={cost43[S:].sum():.2f} 元（固定电价问题3为 14337999.15 元）')

# 指定四天
spec_days = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']
idx = {}
for d in range(N):
    dd = pd.Timestamp(dates[d]).date()
    for sd in spec_days:
        if dd == pd.Timestamp(sd).date():
            idx[sd] = d
print('\n表1/表2/表3 指定日期(波动电价):')
for sd in spec_days:
    d = idx[sd]
    print(f'  {sd}: 4-2购电费={cost42[d]:.2f}, 4-3总费用={cost43[d]:.2f}(计划{rec43[d]["plan_commit"]:.0f}+紧急{rec43[d]["emer_cost"]:.0f})')

print('\n[表4-2a~d] 指定时段计划购电量(kWh)')
for sd in spec_days:
    d = idx[sd]; gd = gE42[d]
    vals = [float(gd[h * 6]) for h in (10, 12, 14, 16, 18, 20)]
    print(f'  {sd}  ' + ' '.join(f'{v:10.2f}' for v in vals)
          + f'  全天={gd.sum():.2f} kWh, 计划费={plan42[d]:.2f} 元, '
            f'紧急量={float(dl42[d].sum() * DT):.2f} kWh, 紧急费={emg42[d]:.2f} 元')

print('\n[表4-2e~h] 指定时段充放电量(kWh)与两端储电量')
for sd in spec_days:
    d = idx[sd]; ce = chE42[d]; de = dchE42[d]
    print(f'  {sd}  0:00储电量={E42[d, 0]:.2f}  24:00储电量={E42[d, -1]:.2f}')
    for a, b in [(0, 4), (4, 8), (8, 12), (12, 16), (16, 20), (20, 24)]:
        k0, k1 = a * 6, b * 6 - 1
        print(f'    {a:02d}:00-{b:02d}:00  充电量={ce[k0:k1 + 1].sum():9.2f}  放电量={de[k0:k1 + 1].sum():9.2f}')

print('\n[表4-2i] 指定日期紧急购电时段与电量')
for sd in spec_days:
    d = idx[sd]; runs = merge_runs(dl42[d])
    print(f'  {sd}: 合计 {float(dl42[d].sum() * DT):.2f} kWh, 共 {len(runs)} 段')
    for a, b, e in runs:
        print(f'    {a}-{b}  {e:.2f} kWh')

# ---------- 出图 ----------
t = np.arange(144) / 6.0
C1, C2, C3, C4, C5, C6 = '#2E5FA3', '#E07B39', '#3BA272', '#C0504D', '#8A6BBE', '#D4A017'
plt.rcParams.update({'figure.dpi': 150, 'savefig.dpi': 300, 'axes.unicode_minus': False})

def style_ax(ax, xlabel=None, ylabel=None):
    ax.spines['top'].set_visible(False); ax.spines['right'].set_visible(False)
    ax.tick_params(direction='in', labelsize=9)
    for s in ['bottom', 'left']:
        ax.spines[s].set_linewidth(0.9)
    if xlabel: ax.set_xlabel(xlabel, fontsize=10)
    if ylabel: ax.set_ylabel(ylabel, fontsize=10)

sd = '2025-09-23'; d0 = idx[sd]

# 图4-1 波动电价 vs 固定电价（典型日）
fig, ax = plt.subplots(figsize=(11.5, 5.2))
ax.plot(t, price_all[d0], color=C4, lw=1.6, label='波动电价(附件4)')
ax.plot(t, price_fixed, color=C1, lw=1.6, ls='--', label='固定电价(附件1)')
style_ax(ax, '小时', '电价 (元/kWh)')
ax.set_xticks(np.arange(0, 25, 6)); ax.set_xlim(-0.3, 24.3)
ax.legend(fontsize=9, loc='upper right', framealpha=0.9)
ax.text(0.01, 0.02, f'(a) {sd} 波动电价与固定电价对比', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图4-1_波动与固定电价对比')

# 图4-2 波动电价四季典型日购电策略（4-2，计划+紧急 叠加）
fig, axes = plt.subplots(1, 4, figsize=(15.6, 4.6))
for ax, sd_, col in zip(axes, spec_days, [C1, C3, C4, C2]):
    d = idx[sd_]
    ax.bar(t, g42[d], width=0.14, color=col, alpha=0.85, edgecolor='black', lw=0.4, label='计划')
    ax.bar(t, dl42[d], bottom=g42[d], width=0.14, color=C5, alpha=0.9, edgecolor='black', lw=0.4, label='紧急')
    ax.set_xticks(np.arange(0, 25, 6)); style_ax(ax, '小时', '购电 (kW)')
    ax.set_ylim(0, max(g42[d].max(), (g42[d] + dl42[d]).max()) * 1.18)
    ax.text(0.03, 0.02, f'({sd_[5:]})', transform=ax.transAxes, fontsize=9, fontweight='bold')
axes[0].legend(fontsize=8, loc='upper left', framealpha=0.9)
save_fig(fig, '图4-2_波动电价四季典型日购电')

# 图4-3 波动电价滚动调整购电量（4-3）
fig, ax = plt.subplots(figsize=(11.5, 5.2))
ax.plot(t, rec43[d0]['g0'], color=C1, lw=1.4, label='0:00计划购电')
ax.plot(t, rec43[d0]['adj'], color=C4, lw=1.4, ls='--', label='滚动调整后购电')
style_ax(ax, '小时', '购电功率 (kW)')
ax.set_xticks(np.arange(0, 25, 6)); ax.set_xlim(-0.3, 24.3)
ax.legend(fontsize=9, loc='upper right', framealpha=0.9)
ax.text(0.01, 0.02, f'(a) {sd} 波动电价下计划与调整购电', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图4-3_波动电价滚动调整购电')

# 图4-4 波动电价SOC轨迹（4-2典型日）
fig, ax = plt.subplots(figsize=(11.5, 5.2))
ax.plot(t, E42[d0, :144], color=C3, lw=1.8, label='4-2 SOC')
ax.plot(t, rec43[d0]['Ef'][:144], color=C4, lw=1.8, ls='--', label='4-3 调整后SOC')
ax.axhline(EMIN, color='#7A7A7A', lw=0.8, ls=':'); ax.axhline(EMAX, color='#7A7A7A', lw=0.8, ls=':')
style_ax(ax, '小时', '储电量 (kWh)')
ax.set_xticks(np.arange(0, 25, 6)); ax.set_ylim(EMIN - 400, EMAX + 400)
ax.legend(fontsize=9, loc='upper right', framealpha=0.9)
ax.text(0.01, 0.02, f'(a) {sd} 波动电价储能SOC轨迹', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图4-4_波动电价SOC轨迹')

# 图4-5 波动电价紧急购电量（4-3，2.1—12.31）
fig, ax = plt.subplots(figsize=(11.5, 5.2))
emer_daily = np.array([rec43[d]['emer_kwh'].sum() for d in range(N)])
ax.bar(dates[S:], emer_daily[S:], width=1.0, color=C4, alpha=0.8, edgecolor='none')
style_ax(ax, '日期', '紧急购电量 (kWh)')
ax.text(0.01, 0.02, '(a) 波动电价下2.1—12.31逐日紧急购电量', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图4-5_波动电价紧急购电量')

# 图4-6 固定vs波动电价费用对比（2.1—12.31）
fig, axes = plt.subplots(1, 2, figsize=(13.6, 5.0))
ax = axes[0]
fixed = [14382892.99, 14337999.15]   # 固定电价下问题2/问题3 的 2.1—12.31 汇总
fluc = [cost42[S:].sum(), cost43[S:].sum()]
x = np.arange(2)
w = 0.35
ax.bar(x - w / 2, [f / 1e4 for f in fixed], w, color=C1, edgecolor='black', lw=0.7, label='固定电价')
ax.bar(x + w / 2, [f / 1e4 for f in fluc], w, color=C4, edgecolor='black', lw=0.7, label='波动电价')
for i in range(2):
    ax.text(i - w / 2, fixed[i] / 1e4 + 10, f'{fixed[i]/1e4:.0f}', ha='center', fontsize=8)
    ax.text(i + w / 2, fluc[i] / 1e4 + 10, f'{fluc[i]/1e4:.0f}', ha='center', fontsize=8)
ax.set_xticks(x); ax.set_xticklabels(['问题2(逐日计划)', '问题3(滚动调整)'], fontsize=10)
style_ax(ax, None, '2.1—12.31购电费 (万元)')
ax.legend(fontsize=9, loc='upper left', framealpha=0.9)
ax.text(0.02, 0.02, '(a) 固定/波动电价2.1—12.31费用', transform=ax.transAxes, fontsize=10, fontweight='bold')
ax = axes[1]
mflu = []
for m in range(1, 12):
    dd = [d for d in range(N) if dates[d].month == m + 1]
    mflu.append(cost43[dd].sum() / 1e4)
ax.bar(np.arange(11), mflu, color=get_colors('tab20', 11), edgecolor='black', lw=0.6)
ax.set_xticks(np.arange(11)); ax.set_xticklabels([f'{m+1}月' for m in range(1, 12)], fontsize=9)
style_ax(ax, '月份', '波动电价4-3月度费用 (万元)')
ax.text(0.02, 0.02, '(b) 波动电价4-3月度费用', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图4-6_固定与波动电价费用对比')

# ---------- 写 result4-2.xlsx ----------
from openpyxl import Workbook
from datetime import time as _midnight
labs = template_labels('0:00-0:10+1')
start = S
wb = Workbook()
ws = wb.active; ws.title = '计划购电量'
ws.append(['日期\\时间'] + labs + ['全天购电量', '全天购电费'])
for d in range(start, N):
    gE = g42[d] * DT
    ws.append([dates[d]] + [round(float(x), 4) for x in gE] + [round(float(gE.sum()), 4), round(float(plan42[d]), 4)])
ws2 = wb.create_sheet('充放电量')
ws2.append(['日期', '时间段', '充电量', '放电量', '时刻', '储电量'])
for d in range(start, N):
    chE = ch42[d] * DT; dchE = dch42[d] * DT
    for bi in range(len(TPL_BIN_LABELS)):
        k0, k1 = bi * 24, bi * 24 + 23
        cq = round(float(chE[k0:k1 + 1].sum()), 4); dq = round(float(dchE[k0:k1 + 1].sum()), 4)
        if bi == 0:
            ws2.append([dates[d], TPL_BIN_LABELS[bi], cq, dq, _midnight(0, 0), round(float(E42[d, 0]), 4)])
        elif bi == 1:
            ws2.append([None, TPL_BIN_LABELS[bi], cq, dq, '24:00', round(float(E42[d, -1]), 4)])
        else:
            ws2.append([None, TPL_BIN_LABELS[bi], cq, dq, None, None])
ws3 = wb.create_sheet('紧急购电量')
ws3.append(['日期', '购电时间段', '购电量'])
for d in range(start, N):
    runs = merge_runs(dl42[d])
    if not runs:
        ws3.append([dates[d], None, None])
        continue
    for i, (a, b, e) in enumerate(runs):
        ws3.append([dates[d] if i == 0 else None, f'{a}-{b}', round(e, 4)])
wb.save(os.path.join(ROOT_OUT, 'result4-2.xlsx'))
wb.save(os.path.join(OUT_DIR, 'result4-2.xlsx'))
print('已保存 result4-2.xlsx → 求解/结果/ 与 问题四/结果/')

# ---------- 写 result4-3.xlsx ----------
wb = Workbook()
ws = wb.active; ws.title = '计划购电量'
ws.append(['日期\\时间'] + labs + ['全天购电量', '全天购电费'])
for d in range(start, N):
    r = rec43[d]; gE = r['g0'] * DT
    ws.append([dates[d]] + [round(float(x), 4) for x in gE] + [round(float(gE.sum()), 4), round(float(r['plan_commit']), 4)])
ws2 = wb.create_sheet('调整购电量')
ws2.append(['日期\\时间'] + labs + ['全天购电量', '全天购电费'])
for d in range(start, N):
    r = rec43[d]; aE = r['adj'] * DT
    ws2.append([dates[d]] + [round(float(x), 4) for x in aE] + [round(float(aE.sum()), 4), round(float(r['purchase_total']), 4)])
ws3 = wb.create_sheet('充放电量')
ws3.append(['日期', '时间段', '充电量', '放电量', '时刻', '储电量'])
for d in range(start, N):
    r = rec43[d]; chE = r['chf'] * DT; dchE = r['dchf'] * DT
    for bi in range(len(TPL_BIN_LABELS)):
        k0, k1 = bi * 24, bi * 24 + 23
        cq = round(float(chE[k0:k1 + 1].sum()), 4); dq = round(float(dchE[k0:k1 + 1].sum()), 4)
        if bi == 0:
            ws3.append([dates[d], TPL_BIN_LABELS[bi], cq, dq, _midnight(0, 0), round(float(r['Ef'][0]), 4)])
        elif bi == 1:
            ws3.append([None, TPL_BIN_LABELS[bi], cq, dq, '24:00', round(float(r['Ef'][-1]), 4)])
        else:
            ws3.append([None, TPL_BIN_LABELS[bi], cq, dq, None, None])
ws4 = wb.create_sheet('紧急购电量')
ws4.append(['日期', '购电时间段', '购电量'])
for d in range(start, N):
    em = rec43[d]['emer_kwh']
    segs = []
    i = 0
    while i < 144:
        if em[i] > 1e-3:
            j = i
            while j < 144 and em[j] > 1e-3:
                j += 1
            segs.append((i, j, em[i:j].sum())); i = j
        else:
            i += 1
    if not segs:
        ws4.append([dates[d], None, None])
    else:
        for si, (a, b, amt) in enumerate(segs):
            period = f'{a*10//60:02d}:{a*10%60:02d}-{b*10//60:02d}:{b*10%60:02d}'
            if si == 0:
                ws4.append([dates[d], period, round(float(amt), 4)])
            else:
                ws4.append([None, period, round(float(amt), 4)])
wb.save(os.path.join(ROOT_OUT, 'result4-3.xlsx'))
wb.save(os.path.join(OUT_DIR, 'result4-3.xlsx'))
print('已保存 result4-3.xlsx → 求解/结果/ 与 问题四/结果/')
