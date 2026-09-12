# 问题三：0/6/12/18时滚动预报 + 调整购电（违约0.5倍/超额1.5倍/紧急5倍）
import sys, os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from _common import *
from _models import *

price_daily, _, _ = read_附件1()
dates, load_all, pv_actual_all = read_附件2()
fc = read_附件3()
print('附件3 预报天数:', len(fc))

def date_key(d):
    return f'{d.year}-{d.month}-{d.day}'

def run_day(d, verbose=False):
    """滚动求解一天，返回 dict"""
    L = load_all[d]
    S_act = pv_actual_all[d]
    key = date_key(dates[d])
    f = fc[key]
    # 各时刻预报 -> 全天10分钟（k0=0..143，未用到的前段由插值外延但不参与计算）
    S0 = forecast_10min(0, f[0], 0, 143)
    S6 = forecast_10min(6, f[6], 0, 143)
    S12 = forecast_10min(12, f[12], 0, 143)
    S18 = forecast_10min(18, f[18], 0, 143)
    # 0:00 计划
    g0, ch0, dch0, E0, c_plan = solve_daily(L, S0, price_daily, e0=carry[d])
    # 6:00 调整
    g6, ch6, dch6, E6, _ = solve_adjust(L, S6, price_daily, g0, e0=float(E0[36]), k0=36)
    # 12:00 调整
    g12, ch12, dch12, E12, _ = solve_adjust(L, S12, price_daily, g0, e0=float(E6[36]), k0=72)
    # 18:00 调整
    g18, ch18, dch18, E18, _ = solve_adjust(L, S18, price_daily, g0, e0=float(E12[36]), k0=108)
    # 拼接最终
    blocks = [0, 36, 72, 108, 144]
    sols = [(g0, ch0, dch0, E0), (g6, ch6, dch6, E6), (g12, ch12, dch12, E12), (g18, ch18, dch18, E18)]
    adj = np.zeros(144); chf = np.zeros(144); dchf = np.zeros(144); Ef = np.empty(145)
    Ef[0] = E0[0]
    for i in range(4):
        a, b = blocks[i], blocks[i + 1]
        g_, c_, d_, E_ = sols[i]
        adj[a:b] = g_[0:b - a]; chf[a:b] = c_[0:b - a]; dchf[a:b] = d_[0:b - a]
        Ef[a + 1:b + 1] = E_[1:b - a + 1]
    # 费用
    plan = g0
    base = np.minimum(plan, adj) * DT * price_daily
    down = np.maximum(plan - adj, 0) * DT * price_daily * 0.5
    up = np.maximum(adj - plan, 0) * DT * price_daily * 1.5
    plan_commit = (plan * DT * price_daily).sum()
    purchase_total = (base + down + up).sum()
    # 紧急购电 = 实际供电缺口：delta_k = max{L - (g_k + S^r_k + d_k - c_k), 0}
    # 其中 g_k 为调整后最终购电、d_k/c_k 为最终放电/充电，S^r_k 为实际光伏。
    gap = np.maximum(L - (adj + S_act + dchf - chf), 0.0)
    emer_kwh = gap * DT
    emer_cost = (emer_kwh * 5.0 * price_daily).sum()
    total = purchase_total + emer_cost
    # 无调整基线：仅执行 0:00 计划，缺口按相同定义核算
    gap0 = np.maximum(L - (g0 + S_act + dch0 - ch0), 0.0)
    emer0_kwh = gap0 * DT
    total0 = plan_commit + (emer0_kwh * 5.0 * price_daily).sum()
    return dict(g0=g0, adj=adj, chf=chf, dchf=dchf, Ef=Ef, S0=S0, S6=S6, S12=S12, S18=S18,
                S_act=S_act, plan_commit=plan_commit, purchase_total=purchase_total,
                emer_kwh=emer_kwh, emer_cost=emer_cost, total=total, total0=total0)

def _fc_for(T, key, S_act, mode):
    """返回按时刻T发布的10分钟光伏预报（全144）。
    T=0 恒用 0:00 真实预报；mode='perfect' 时 T>0 全部用实际光伏（完全预见上界）；
    mode='real' 时附件3有预报的时刻用真实预报，其余时刻用实际光伏近似。"""
    f = fc[key]
    if T == 0:
        return forecast_10min(0, f[0], 0, 143)
    if mode == 'perfect' or T not in f:
        return S_act
    return forecast_10min(T, f[T], 0, 143)

def run_day_times(d, times, e0_start, mode='real'):
    """按给定调整时刻序列 times(小时,含0) 滚动求解一天，返回(当天总费用, 次日初始储电量)。
    计划购电恒由 0:00 制定；后续每个时刻在剩余时段 [T,24:00] 重优化调整购电，
    购电费用=min(plan,adj)+0.5*max(plan-adj,0)+1.5*max(adj-plan,0)，缺口按5倍紧急购电。"""
    L = load_all[d]
    S_act = pv_actual_all[d]
    key = date_key(dates[d])
    S_list = [_fc_for(T, key, S_act, mode) for T in times]
    g0, ch0, dch0, E0, _ = solve_daily(L, S_list[0], price_daily, e0=e0_start)
    sols = [(g0, ch0, dch0, E0)]
    prev_k0, prev_E = 0, E0
    for i in range(1, len(times)):
        T = times[i]
        k0 = T * 6
        e0 = float(prev_E[k0 - prev_k0])
        adj, ch, dch, E, _ = solve_adjust(L, S_list[i], price_daily, g0, e0=e0, k0=k0)
        sols.append((adj, ch, dch, E))
        prev_k0, prev_E = k0, E
    ks = [0] + [T * 6 for T in times[1:]] + [144]
    adj = np.zeros(144); chf = np.zeros(144); dchf = np.zeros(144); Ef = np.empty(145)
    Ef[0] = E0[0]
    for i in range(len(times)):
        a, b = ks[i], ks[i + 1]
        g_, c_, d_, E_ = sols[i]
        adj[a:b] = g_[0:b - a]; chf[a:b] = c_[0:b - a]; dchf[a:b] = d_[0:b - a]
        Ef[a + 1:b + 1] = E_[1:b - a + 1]
    plan = g0
    base = np.minimum(plan, adj) * DT * price_daily
    down = np.maximum(plan - adj, 0) * DT * price_daily * 0.5
    up = np.maximum(adj - plan, 0) * DT * price_daily * 1.5
    purchase_total = (base + down + up).sum()
    gap = np.maximum(L - (adj + S_act + dchf - chf), 0.0)
    emer_cost = (gap * DT * 5.0 * price_daily).sum()
    return purchase_total + emer_cost, float(Ef[-1])

def run_year(times, mode='real', start=0):
    """全年滚动求总费用（carry 储电量逐日传递）。
    start 为汇总起始日索引（附件5要求仅统计 2.1—12.31，1月作为预热期，故 start=31）。"""
    carry2 = [E0_INIT] * N
    tot = 0.0
    for d in range(N):
        t, e_end = run_day_times(d, times, carry2[d], mode)
        if d >= start:
            tot += t
        if d + 1 < N:
            carry2[d + 1] = e_end
    return tot

# 全年滚动（carry：储电量传递）
N = len(dates)
S = 31   # 2025-02-01 索引；附件5仅要求保存 2.1—12.31，1月作为预热期，不纳入汇总统计
carry = [E0_INIT] * N
rec = [None] * N
totals = np.zeros(N); totals0 = np.zeros(N)
for d in range(N):
    r = run_day(d)
    rec[d] = r
    totals[d] = r['total']; totals0[d] = r['total0']
    if d + 1 < N:
        carry[d + 1] = float(r['Ef'][-1])
print(f'全年(365天,含1月预热)含调整总费用={totals.sum():.2f} 元, 无调整={totals0.sum():.2f} 元')
print(f'[汇总口径] 2.1—12.31 含调整={totals[S:].sum():.2f} 元, 无调整={totals0[S:].sum():.2f} 元, '
      f'节省={totals0[S:].sum()-totals[S:].sum():.2f} 元, 降幅={100*(totals0[S:].sum()-totals[S:].sum())/totals0[S:].sum():.3f}%')

# 指定四天详细结果
spec_days = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']
idx = {}
for d in range(N):
    dd = pd.Timestamp(dates[d]).date()
    for sd in spec_days:
        if dd == pd.Timestamp(sd).date():
            idx[sd] = d
print('\n表1/表2/表3 指定日期:')
for sd in spec_days:
    d = idx[sd]; r = rec[d]
    gE = r['g0'] * DT; aE = r['adj'] * DT
    print(f'  {sd}: 计划购电量={gE.sum():.1f}, 调整购电量={aE.sum():.1f}, '
          f'总费用={r["total"]:.2f}(计划{r["plan_commit"]:.0f}+调整{r["purchase_total"]-r["plan_commit"]:.0f}+紧急{r["emer_cost"]:.0f})')
    for hh in [10, 12, 14, 16, 18, 20]:
        k = hh * 6
        print(f'     {hh:02d}:00-{hh:02d}:10  计划={gE[k]:.2f}  调整={aE[k]:.2f}')

# ---------- 出图 ----------
t = np.arange(144) / 6.0
C1, C2, C3, C4, C5, C6 = '#2E5FA3', '#E07B39', '#3BA272', '#C0504D', '#8A6BBE', '#D4A017'
plt.rcParams.update({'figure.dpi': 150, 'savefig.dpi': 300, 'axes.unicode_minus': False})

def style_ax(ax, xlabel=None, ylabel=None):
    ax.spines['top'].set_visible(False); ax.spines['right'].set_visible(False)
    ax.tick_params(direction='in', labelsize=9)
    for s in ['bottom', 'left']:
        ax.spines[s].set_linewidth(0.9)
    if xlabel: ax.set_xlabel(xlabel, fontsize=10)
    if ylabel: ax.set_ylabel(ylabel, fontsize=10)

sd = '2025-09-23'; d0 = idx[sd]; r0 = rec[d0]

# 图3-1 各时刻预报与实际光伏
fig, ax = plt.subplots(figsize=(11.5, 5.2))
ax.plot(t, r0['S_act'], color='black', lw=2.0, label='实际光伏')
ax.plot(t, r0['S0'], color=C1, lw=1.2, label='0:00预报')
ax.plot(t[36:], r0['S6'][36:], color=C2, lw=1.2, label='6:00预报')
ax.plot(t[72:], r0['S12'][72:], color=C3, lw=1.2, label='12:00预报')
ax.plot(t[108:], r0['S18'][108:], color=C4, lw=1.2, label='18:00预报')
style_ax(ax, '小时', '光伏功率 (kW)')
ax.set_xticks(np.arange(0, 25, 6)); ax.set_xlim(-0.3, 24.3)
ax.legend(fontsize=9, loc='upper right', framealpha=0.9)
ax.text(0.01, 0.02, f'(a) {sd} 各时刻预报与实际光伏', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图3-1_预报与实际光伏对比')

# 图3-2 计划与最终调整购电量对比
fig, ax = plt.subplots(figsize=(11.5, 5.2))
ax.plot(t, r0['g0'], color=C1, lw=1.4, label='0:00计划购电')
ax.plot(t, r0['adj'], color=C4, lw=1.4, ls='--', label='滚动调整后购电')
style_ax(ax, '小时', '购电功率 (kW)')
ax.set_xticks(np.arange(0, 25, 6)); ax.set_xlim(-0.3, 24.3)
ax.legend(fontsize=9, loc='upper right', framealpha=0.9)
ax.text(0.01, 0.02, f'(a) {sd} 计划购电与调整购电对比', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图3-2_计划与调整购电量对比')

# 图3-3 2.1—12.31 逐日紧急购电量
fig, ax = plt.subplots(figsize=(11.5, 5.2))
emer_daily = np.array([rec[d]['emer_kwh'].sum() for d in range(N)])
ax.bar(dates[S:], emer_daily[S:], width=1.0, color=C4, alpha=0.8, edgecolor='none')
style_ax(ax, '日期', '紧急购电量 (kWh)')
ax.text(0.01, 0.02, '(a) 2.1—12.31逐日紧急购电量', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图3-3_全年紧急购电量')

# 图3-4 有无调整费用对比 + 月度节省（2—12月）
fig, axes = plt.subplots(1, 2, figsize=(13.6, 5.0))
ax = axes[0]
ax.bar([0, 1], [totals0[S:].sum() / 1e4, totals[S:].sum() / 1e4], width=0.5,
       color=[C1, C3], edgecolor='black', lw=0.8)
for i, v in enumerate([totals0[S:].sum() / 1e4, totals[S:].sum() / 1e4]):
    ax.text(i, v + 15, f'{v:.1f}万元', ha='center', fontsize=10)
ax.set_xticks([0, 1]); ax.set_xticklabels(['仅0:00计划', '滚动调整'], fontsize=10)
style_ax(ax, None, '2.1—12.31购电费 (万元)')
ax.text(0.03, 0.02, '(a) 有无调整2.1—12.31费用', transform=ax.transAxes, fontsize=10, fontweight='bold')
ax = axes[1]
msave = []
for m in range(1, 12):
    dd = [d for d in range(N) if dates[d].month == m + 1]
    msave.append((totals0[dd].sum() - totals[dd].sum()) / 1e4)
ax.bar(np.arange(11), msave, color=get_colors('tab20', 11), edgecolor='black', lw=0.6)
ax.set_xticks(np.arange(11)); ax.set_xticklabels([f'{m+1}月' for m in range(1, 12)], fontsize=9)
style_ax(ax, '月份', '调整节省费用 (万元)')
ax.text(0.02, 0.02, '(b) 月度调整节省', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图3-4_有无调整费用对比')

# 图3-5 典型日SOC：计划 vs 调整后
fig, ax = plt.subplots(figsize=(11.5, 5.2))
# 重新求解0:00计划SOC（用于对比）
g0_, ch0_, dch0_, E0p_, _ = solve_daily(load_all[d0], rec[d0]['S0'], price_daily, e0=carry[d0])
ax.plot(t, E0p_[:144], color=C1, lw=1.3, ls='--', label='0:00计划SOC')
ax.plot(t, rec[d0]['Ef'][:144], color=C4, lw=1.8, label='滚动调整后SOC')
ax.axhline(EMIN, color='#7A7A7A', lw=0.8, ls=':'); ax.axhline(EMAX, color='#7A7A7A', lw=0.8, ls=':')
style_ax(ax, '小时', '储电量 (kWh)')
ax.set_xticks(np.arange(0, 25, 6)); ax.set_ylim(EMIN - 400, EMAX + 400)
ax.legend(fontsize=9, loc='upper right', framealpha=0.9)
ax.text(0.01, 0.02, f'(a) {sd} 储电量：计划 vs 调整后', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图3-5_计划与调整后SOC对比')

# 图3-6 费用构成（2—12月，月度堆叠）
fig, ax = plt.subplots(figsize=(11.5, 5.2))
plan_m = []; adj_m = []; emer_m = []
for m in range(1, 12):
    dd = [d for d in range(N) if dates[d].month == m + 1]
    plan_m.append(sum(rec[d]['plan_commit'] for d in dd) / 1e4)
    adj_m.append(sum(rec[d]['purchase_total'] - rec[d]['plan_commit'] for d in dd) / 1e4)
    emer_m.append(sum(rec[d]['emer_cost'] for d in dd) / 1e4)
x = np.arange(11)
ax.bar(x, plan_m, width=0.6, color=C1, edgecolor='black', lw=0.5, label='计划购电费')
ax.bar(x, adj_m, width=0.6, bottom=plan_m, color=C2, edgecolor='black', lw=0.5, label='调整相关费')
ax.bar(x, emer_m, width=0.6, bottom=np.array(plan_m) + np.array(adj_m), color=C4, edgecolor='black', lw=0.5, label='紧急购电费')
ax.set_xticks(x); ax.set_xticklabels([f'{m+1}月' for m in range(1, 12)], fontsize=9)
style_ax(ax, '月份', '月度费用 (万元)')
ax.legend(fontsize=9, loc='upper left', framealpha=0.9)
ax.text(0.02, 0.02, '(a) 2—12月费用构成(月度堆叠)', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图3-6_全年费用构成')

# ---------- 写 result3.xlsx ----------
from openpyxl import Workbook
from datetime import time as _midnight
out = os.path.join(ROOT_OUT, 'result3.xlsx')
wb = Workbook()
labs = template_labels('0:00-0:10+1')
ws = wb.active; ws.title = '计划购电量'
ws.append(['日期\\时间'] + labs + ['全天购电量', '全天购电费'])
start = S
for d in range(start, N):
    r = rec[d]
    gE = r['g0'] * DT
    ws.append([dates[d]] + [round(float(x), 4) for x in gE] + [round(float(gE.sum()), 4), round(float(r['plan_commit']), 4)])
ws2 = wb.create_sheet('调整购电量')
ws2.append(['日期\\时间'] + labs + ['全天购电量', '全天购电费'])
for d in range(start, N):
    r = rec[d]
    aE = r['adj'] * DT
    ws2.append([dates[d]] + [round(float(x), 4) for x in aE] + [round(float(aE.sum()), 4), round(float(r['purchase_total']), 4)])
ws3 = wb.create_sheet('充放电量')
ws3.append(['日期', '时间段', '充电量', '放电量', '时刻', '储电量'])
for d in range(start, N):
    r = rec[d]
    chE = r['chf'] * DT; dchE = r['dchf'] * DT
    for bi in range(len(TPL_BIN_LABELS)):
        k0, k1 = bi * 24, bi * 24 + 23
        cq = round(float(chE[k0:k1 + 1].sum()), 4); dq = round(float(dchE[k0:k1 + 1].sum()), 4)
        if bi == 0:
            ws3.append([dates[d], TPL_BIN_LABELS[bi], cq, dq, _midnight(0, 0), round(float(r['Ef'][0]), 4)])
        elif bi == 1:
            ws3.append([None, TPL_BIN_LABELS[bi], cq, dq, '24:00', round(float(r['Ef'][-1]), 4)])
        else:
            ws3.append([None, TPL_BIN_LABELS[bi], cq, dq, None, None])
ws4 = wb.create_sheet('紧急购电量')
ws4.append(['日期', '紧急购电时间段', '紧急购电量'])
for d in range(start, N):
    r = rec[d]
    em = r['emer_kwh']
    segs = []
    i = 0
    while i < 144:
        if em[i] > 1e-3:
            j = i
            while j < 144 and em[j] > 1e-3:
                j += 1
            segs.append((i, j, em[i:j].sum()))
            i = j
        else:
            i += 1
    if not segs:
        ws4.append([dates[d], None, None])
    else:
        for si, (a, b, amt) in enumerate(segs):
            period = f'{a*10//60:02d}:{a*10%60:02d}-{b*10//60:02d}:{b*10%60:02d}'
            if si == 0:
                ws4.append([dates[d], period, round(float(amt), 4)])
            else:
                ws4.append([None, period, round(float(amt), 4)])
wb.save(out)
wb.save(os.path.join(OUT_DIR, 'result3.xlsx'))
print(f'已保存 result3.xlsx（{N - start} 天）→ 支撑材料/result/')

# ---------- 边际收益分析：是否需要引入更多时刻的预报 ----------
print('\n========== 其他时刻预报的边际收益分析 ==========')
base0 = totals0[S:].sum()                  # 仅0:00（无任何调整），2.1—12.31 口径
cur = totals[S:].sum()                     # 当前方案 0/6/12/18（真实预报），2.1—12.31 口径
# 真实可用预报逐次累加（频率与精度均为题目给定水平）
s_real = [
    ('仅0:00',    base0),
    ('0/6',       run_year([0, 6], start=S)),
    ('0/6/12',    run_year([0, 6, 12], start=S)),
    ('0/6/12/18', cur),
]
# 完全预见（相同预报质量下）的频率效应：更密/更多时刻是否有额外收益
s_perf = [
    ('每6h(完全预见)',      run_year([0, 6, 12, 18], mode='perfect', start=S)),
    ('每6h+9/15(完全预见)', run_year([0, 6, 9, 12, 15, 18], mode='perfect', start=S)),
    ('每3h(完全预见)',      run_year([0, 3, 6, 9, 12, 15, 18, 21], mode='perfect', start=S)),
    ('每小时(完全预见)',    run_year(list(range(24)), mode='perfect', start=S)),
]
for nm, c in s_real + s_perf:
    print(f'  {nm}: {c:,.2f} 元, 较仅0:00节省 {base0-c:,.2f} 元 ({100*(base0-c)/base0:.3f}%)')
r = {nm: c for nm, c in s_real + s_perf}
m6 = base0 - r['0/6']
m12 = r['0/6'] - r['0/6/12']
m18 = r['0/6/12'] - r['0/6/12/18']
m915 = r['每6h(完全预见)'] - r['每6h+9/15(完全预见)']
m3h = r['每6h(完全预见)'] - r['每3h(完全预见)']
m1h = r['每3h(完全预见)'] - r['每小时(完全预见)']
acc = cur - r['每6h(完全预见)']
print(f'  真实预报边际节省: 6:00={m6:,.2f} 元, 12:00={m12:,.2f} 元, 18:00={m18:,.2f} 元')
print(f'  完全预见下频率增量: +9/15={m915:,.2f} 元, 6h→3h={m3h:,.2f} 元, 3h→1h={m1h:,.2f} 元')
print(f'  精度效应(6h频率下): 真实→完全预见 再省 {acc:,.2f} 元')
print(f'  当前方案捕获比 = {100*(base0-cur)/(base0-r["每小时(完全预见)"]):.2f}% (相对每小时完全预见上界)')

# 图3-7 其他时刻预报的边际收益
fig, axes = plt.subplots(1, 2, figsize=(13.6, 5.0))
# (a) 全年总费用
ax = axes[0]
lab_a = ['仅0:00', '0/6', '0/6/12', '0/6/12/18', '每6h\n完全预见', '每3h\n完全预见', '每小时\n完全预见']
val_a = [s_real[i][1] for i in range(4)] + [s_perf[0][1], s_perf[2][1], s_perf[3][1]]
col_a = [C1, C1, C1, C3] + [C6, C6, C6]
bars_a = [v / 1e4 for v in val_a]
ax.bar(range(7), bars_a, width=0.6, color=col_a, edgecolor='black', lw=0.6)
for i, v in enumerate(bars_a):
    ax.text(i, v + 6, f'{v:.1f}', ha='center', fontsize=8)
ax.set_xticks(range(7)); ax.set_xticklabels(lab_a, fontsize=8.5)
# 动态 y 轴范围：完全预见方案费用远低于真实预报（约1220万元，低于原固定下限1380），
# 固定 ylim 会把这三根柱子裁掉，导致"空白"。改为按数据自动缩放并留出标签空间。
lo_a, hi_a = min(bars_a), max(bars_a)
pad_a = max((hi_a - lo_a) * 0.18, 6.0)
ax.set_ylim(lo_a - pad_a, hi_a + pad_a)
style_ax(ax, None, '2.1—12.31总费用 (万元)')
ax.text(0.02, 0.02, '(a) 不同调整时刻方案2.1—12.31总费用', transform=ax.transAxes, fontsize=10, fontweight='bold')
# (b) 边际节省
ax = axes[1]
lab_b = ['6:00', '12:00', '18:00', '+9:00/\n15:00', '6h→3h', '3h→1h']
val_b = [m6, m12, m18, m915, m3h, m1h]
col_b = [C3, C3, C3, C6, C6, C6]
ax.bar(range(6), [v / 1e4 for v in val_b], width=0.6, color=col_b, edgecolor='black', lw=0.6)
for i, v in enumerate(val_b):
    ax.text(i, v / 1e4 + 0.8, f'{v/1e4:.2f}', ha='center', fontsize=8)
ax.set_xticks(range(6)); ax.set_xticklabels(lab_b, fontsize=8.5)
style_ax(ax, None, '边际节省 (万元)')
import matplotlib.patches as mpatches
ax.legend(handles=[mpatches.Patch(color=C3, label='真实预报时刻'),
                   mpatches.Patch(color=C6, label='完全预见下加密时刻')],
          fontsize=8, loc='upper right', framealpha=0.9)
ax.text(0.02, 0.02, '(b) 各新增时刻的边际节省', transform=ax.transAxes, fontsize=10, fontweight='bold')
save_fig(fig, '图3-7_其他时刻预报边际收益')
