# -*- coding: utf-8 -*-
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.join('C题', 'thesis', '参考文献.docx')

SEC1 = '一、中文期刊文献'
SEC2 = '二、专著与教材'
SEC3 = '三、英文文献'

SECTIONS = [
    (SEC1, [
        ('[1] 鲁宗相, 王彩霞, 闵勇, 等. 微电网研究综述[J]. 电力系统自动化, 2007, 31(19): 100-107.',
         '微电网概念、结构与运行控制方式的权威综述, 可在引言中用来界定微电网定义与研究背景。'),
        ('[2] 杨新法, 苏剑, 吕志鹏, 等. 微电网技术综述[J]. 中国电机工程学报, 2014, 34(1): 57-70.',
         '系统梳理微电网建模、控制与能量管理关键技术, 可支撑本文光储微电网系统架构的描述。'),
        ('[3] 王成山, 武震, 李鹏. 微电网关键技术研究[J]. 电工技术学报, 2014, 29(2): 1-12.',
         '微电网关键技术高被引文献, 可引用其对微电网优化运行与能量管理研究方向的归纳。'),
        ('[4] 丁明, 张颖媛, 茆美琴. 微网研究中的关键技术[J]. 电网技术, 2009, 33(11): 6-11.',
         '微网关键技术早期经典文献, 可用于说明微网经济运行问题的提出背景与技术脉络。'),
        ('[5] 周林, 黄勇, 郭珂, 等. 微电网储能技术研究综述[J]. 电力系统保护与控制, 2011, 39(7): 147-152.',
         '综述储能在微电网中的作用, 可支撑"储能平抑光伏波动、削峰填谷"的建模动机。'),
        ('[6] 李霞林, 郭力, 王成山, 等. 直流微电网关键技术研究综述[J]. 中国电机工程学报, 2016, 36(1): 2-17.',
         '直流微电网关键技术综述, 可用于比较交流/直流微电网结构及变换器损耗建模差异。'),
        ('[7] 吴雄, 王秀丽, 刘世民, 等. 微电网能量管理系统研究综述[J]. 电力自动化设备, 2014, 34(10): 7-14.',
         '微电网能量管理系统(EMS)综述, 可引用其能量管理功能架构作为本文调度框架的依据。'),
        ('[8] 桑丙玉, 王德顺, 杨波, 等. 平滑新能源输出波动的储能优化配置方法[J]. 中国电机工程学报, 2014, 34(22): 3700-3706.',
         '用储能平滑新能源出力波动的优化配置方法, 可支撑储能功率/容量约束的建模思路。'),
        ('[9] 张钦, 王锡凡, 王建学, 等. 电力市场下需求响应研究综述[J]. 电力系统自动化, 2008, 32(3): 97-106.',
         '电力市场下需求响应研究综述, 可支撑分时电价机制与负荷响应潜力的分析。'),
        ('[10] 刘卫亮, 刘长良, 林永君, 等. 计及雾霾影响因素的光伏发电超短期功率预测[J]. 中国电机工程学报, 2018, 38(14): 4086-4095.',
         '光伏超短期功率预测的代表性研究, 可支撑光伏出力预测建模与预测误差处理。'),
    ]),
    (SEC2, [
        ('[11] 《运筹学》教材编写组. 运筹学(第5版)[M]. 北京: 清华大学出版社, 2021.',
         '线性规划建模与单纯形法求解原理的经典教材依据, 可在模型求解方法处引用。'),
        ('[12] 姜启源, 谢金星, 叶俊. 数学模型(第五版)[M]. 北京: 高等教育出版社, 2018.',
         '数学建模方法论与优化模型章节的通用参考书, 可作为整体建模流程的方法依据。'),
        ('[13] 盛骤, 谢式千, 潘承毅. 概率论与数理统计(第四版)[M]. 北京: 高等教育出版社, 2008.',
         '概率分布、置信区间与假设检验方法的经典教材, 可支撑数据处理与误差统计检验。'),
    ]),
    (SEC3, [
        ('[14] Virtanen P, Gommers R, Oliphant T E, et al. SciPy 1.0: fundamental algorithms for scientific computing in Python[J]. Nature Methods, 2020, 17(3): 261-272.',
         '本文线性规划求解所依赖的 SciPy 科学计算库的权威出处, 可在求解工具说明处引用。'),
        ('[15] Palma-Behnke R, Benavides C, Lanas F, et al. A microgrid energy management system based on the rolling horizon strategy[J]. IEEE Transactions on Smart Grid, 2013, 4(2): 996-1006.',
         '基于滚动时域(rolling horizon)的微电网能量管理系统, 是本文滚动调度策略的代表性依据。'),
        ('[16] Parisio A, Rikos E, Glielmo L. A model predictive control approach to microgrid operation optimization[J]. IEEE Transactions on Control Systems Technology, 2014, 22(5): 1813-1827.',
         '微电网运行优化的模型预测控制(MPC)方法, 可支撑"预测-优化-校正"调度框架及对比实验。'),
        ('[17] Huangfu Q, Hall J A J. Parallelizing the dual revised simplex method[J]. Mathematical Programming Computation, 2018, 10(1): 119-142.',
         '对偶单纯形法并行化的代表性工作, 可在大规模线性规划求解算法说明处引用。'),
    ]),
]


def style_run(run, size, bold=False, gray=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(size)
    run.font.bold = bold
    if gray:
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


doc = Document()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_run(p.add_run('参考文献'), 16, bold=True)

for sec_title, items in SECTIONS:
    p = doc.add_paragraph()
    style_run(p.add_run(sec_title), 12, bold=True)
    for entry, point in items:
        p = doc.add_paragraph()
        style_run(p.add_run(entry), 12)
        p = doc.add_paragraph()
        style_run(p.add_run('【可引用点】' + point), 9, gray=True)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)

# ---- verify: reopen and count entries per section ----
doc2 = Document(OUT)
counts = {SEC1: 0, SEC2: 0, SEC3: 0}
cur = None
for para in doc2.paragraphs:
    t = para.text.strip()
    if t in counts:
        cur = t
    elif cur and re.match(r'^\[\d+\]', t):
        counts[cur] += 1
total = sum(counts.values())
print('sec1=%d sec2=%d sec3=%d total=%d' % (counts[SEC1], counts[SEC2], counts[SEC3], total))
