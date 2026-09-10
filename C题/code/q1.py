# -*- coding: utf-8 -*-
"""问题1：单日周期 LP 购电策略"""
import numpy as np
import pandas as pd
from pathlib import Path
from scipy.optimize import linprog
from scipy.sparse import lil_matrix, vstack
import shutil
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent

DT = 1.0 / 6.0          # 10 min in hours
T = 144                 # periods per day
PMAX = 5000.0 * DT      # max charge/discharge energy per period (kWh)
ETA = 0.9
SOC_LO, SOC_HI = 1200.0, 10800.0

# ---------- data ----------
a1 = pd.read_excel(BASE / '附件/附件1.xlsx')
price = a1['电价'].to_numpy(float)          # 144
load  = a1['小区负载'].to_numpy(float) * DT # kWh per period
pv    = a1['光伏发电预测功率'].to_numpy(float) * DT

# ---------- LP ----------
# variables: buy[0:144], ch[144:288], dis[288:432], soc[432:577]  (soc has 145 entries)
N = 3 * T + T + 1
ib, ic, id_, is_ = 0, T, 2 * T, 3 * T

c = np.zeros(N)
c[ib:ib + T] = price

# equality: soc dynamics  soc[t+1] - soc[t] - ETA*ch[t] + dis[t]/ETA = 0
Aeq = lil_matrix((T + 1, N))
beq = np.zeros(T + 1)
for t in range(T):
    Aeq[t, is_ + t + 1] = 1.0
    Aeq[t, is_ + t] = -1.0
    Aeq[t, ic + t] = -ETA
    Aeq[t, id_ + t] = 1.0 / ETA
# daily cycle: soc[144] - soc[0] = 0
Aeq[T, is_ + T] = 1.0
Aeq[T, is_] = -1.0

# inequality: -buy - dis + ch <= pv - load   (i.e. buy+pv+dis >= load+ch)
Aub = lil_matrix((T, N))
bub = pv - load
for t in range(T):
    Aub[t, ib + t] = -1.0
    Aub[t, id_ + t] = -1.0
    Aub[t, ic + t] = 1.0

bounds = ([(0, None)] * T + [(0, PMAX)] * T + [(0, PMAX)] * T +
          [(SOC_LO, SOC_HI)] * (T + 1))

res = linprog(c, A_ub=Aub.tocsr(), b_ub=bub, A_eq=Aeq.tocsr(), b_eq=beq,
              bounds=bounds, method='highs')
assert res.status == 0, res.message
x = res.x
buy = x[ib:ib + T]
ch = x[ic:ic + T]
dis = x[id_:id_ + T]
soc = x[is_:is_ + T + 1]

# ---------- checks ----------
bal = buy + pv + dis - load - ch
assert bal.min() > -1e-6
assert soc.min() >= SOC_LO - 1e-6 and soc.max() <= SOC_HI + 1e-6
assert abs(soc[-1] - soc[0]) < 1e-6
simul = np.minimum(ch, dis).max()
print('max simultaneous ch/dis (should ~0):', simul)

total_buy = buy.sum()
total_cost = float(price @ buy)
print('全天购电量 kWh: %.4f' % total_buy)
print('全天购电费 元 : %.4f' % total_cost)
print('0:00 储电量: %.4f   24:00 储电量: %.4f' % (soc[0], soc[-1]))

# 表1 specified periods (period starting at given clock time)
slots = {'10:00-10:10': 60, '12:00-12:10': 72, '14:00-14:10': 84,
         '16:00-16:10': 96, '18:00-18:10': 108, '20:00-20:10': 120}
print('\n表1 指定时段购电量:')
for k, i in slots.items():
    print('  %s : %.4f' % (k, buy[i]))

# 表2 4-hour blocks
print('\n表2 4小时时段充/放电量:')
blocks = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
names = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00',
         '16:00-20:00', '20:00-24:00']
for n, (a, b) in zip(names, blocks):
    print('  %-12s 充电 %.4f  放电 %.4f' % (n, ch[a:b].sum(), dis[a:b].sum()))

# ---------- write result1.xlsx ----------
OUT1 = BASE / 'result1.xlsx'
shutil.copy(BASE / '附件/附件5/result1.xlsx', OUT1)
wb = load_workbook(OUT1)
ws = wb['计划购电量']
for t in range(T):
    ws.cell(row=t + 2, column=2, value=round(float(buy[t]), 4))
ws2 = wb['充放电量']
for r, (a, b) in enumerate(blocks):
    ws2.cell(row=r + 2, column=2, value=round(float(ch[a:b].sum()), 4))
    ws2.cell(row=r + 2, column=3, value=round(float(dis[a:b].sum()), 4))
ws2.cell(row=2, column=5, value=round(float(soc[0]), 4))    # 0:00
ws2.cell(row=3, column=5, value=round(float(soc[-1]), 4))   # 24:00
wb.save(OUT1)
print('\nresult1.xlsx 已生成')
