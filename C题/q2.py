# -*- coding: utf-8 -*-
"""问题2：全年逐日 LP（SOC 跨天连续传递），输出 result2.xlsx"""
import numpy as np
import pandas as pd
from scipy.optimize import linprog
from scipy.sparse import lil_matrix
import shutil, io, datetime
from openpyxl import load_workbook

DT = 1.0 / 6.0
T = 144
PMAX = 5000.0 * DT
ETA = 0.9
SOC_LO, SOC_HI = 1200.0, 10800.0
EMERG = 5.0

log = io.open('q2_log.txt', 'w', encoding='utf-8')
def pr(*a):
    s = ' '.join(str(x) for x in a)
    log.write(s + '\n')
    print(s.encode('gbk', 'replace').decode('gbk'))

# ---------- data ----------
a1 = pd.read_excel('附件/附件1.xlsx')
price = a1['电价'].to_numpy(float)                      # 144, same every day
load_df = pd.read_excel('附件/附件2.xlsx', sheet_name='小区负载')
pv_df   = pd.read_excel('附件/附件2.xlsx', sheet_name='光伏发电实际功率')
dates = pd.to_datetime(load_df.iloc[:, 0])
LOAD = load_df.iloc[:, 1:].to_numpy(float) * DT         # 365 x 144 kWh
PV   = pv_df.iloc[:, 1:].to_numpy(float) * DT
NDAYS = LOAD.shape[0]
assert NDAYS == 365

# ---------- LP template ----------
# vars: buy[0:T], ch[T:2T], dis[2T:3T], e[3T:4T], soc[4T:5T+1]
N = 4 * T + T + 1
ib, ic, id_, ie, is_ = 0, T, 2 * T, 3 * T, 4 * T

c = np.zeros(N)
c[ib:ib + T] = price
c[ie:ie + T] = EMERG * price

Aeq = lil_matrix((T, N))
beq = np.zeros(T)
for t in range(T):
    Aeq[t, is_ + t + 1] = 1.0
    Aeq[t, is_ + t] = -1.0
    Aeq[t, ic + t] = -ETA
    Aeq[t, id_ + t] = 1.0 / ETA
Aeq = Aeq.tocsr()

Aub = lil_matrix((T, N))
for t in range(T):
    Aub[t, ib + t] = -1.0
    Aub[t, id_ + t] = -1.0
    Aub[t, ie + t] = -1.0
    Aub[t, ic + t] = 1.0
Aub = Aub.tocsr()

base_bounds = ([(0, None)] * T + [(0, PMAX)] * T + [(0, PMAX)] * T +
               [(0, None)] * T + [(SOC_LO, SOC_HI)] * (T + 1))

# ---------- rolling solve ----------
soc0 = 6000.0
results = {}
for d in range(NDAYS):
    bub = PV[d] - LOAD[d]
    bounds = list(base_bounds)
    bounds[4 * T] = (soc0, soc0)          # fix soc[0]
    res = linprog(c, A_ub=Aub, b_ub=bub, A_eq=Aeq, b_eq=beq,
                  bounds=bounds, method='highs')
    assert res.status == 0, (d, res.message)
    x = res.x
    buy, ch, dis = x[ib:ib+T], x[ic:ic+T], x[id_:id_+T]
    e, soc = x[ie:ie+T], x[is_:is_+T+1]
    # checks
    assert (buy + PV[d] + dis + e - LOAD[d] - ch).min() > -1e-5
    assert soc.min() >= SOC_LO - 1e-6 and soc.max() <= SOC_HI + 1e-6
    results[d] = dict(buy=buy, ch=ch, dis=dis, e=e, soc=soc,
                      cost=float(price @ buy + EMERG * price @ e))
    soc0 = float(soc[-1])

# ---------- summary ----------
all_e = np.array([results[d]['e'] for d in range(NDAYS)])
pr('全年紧急购电总量 kWh: %.6f' % all_e.sum())
pr('紧急购电非零天数: %d' % (all_e.sum(axis=1) > 1e-6).sum())
costs = np.array([results[d]['cost'] for d in range(NDAYS)])
pr('全年总购电费 元: %.2f' % costs.sum())
pr('1月31天总费用: %.2f  2.1-12.31总费用: %.2f' % (costs[:31].sum(), costs[31:].sum()))
pr('2025-02-01 0:00 SOC: %.4f' % results[31]['soc'][0])
pr('2025-12-31 24:00 SOC: %.4f' % results[364]['soc'][-1])

# 4 specified dates
slots = [('10:00-10:10', 60), ('12:00-12:10', 72), ('14:00-14:10', 84),
         ('16:00-16:10', 96), ('18:00-18:10', 108), ('20:00-20:10', 120)]
blocks = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
bnames = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00',
          '16:00-20:00', '20:00-24:00']
for ds in ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']:
    d = int(np.where(dates == pd.Timestamp(ds))[0][0])
    r = results[d]
    pr('\n===== %s (day idx %d) =====' % (ds, d))
    pr('全天购电量: %.4f kWh   全天购电费: %.4f 元' % (r['buy'].sum(), r['cost']))
    for k, i in slots:
        pr('  %s : %.4f' % (k, r['buy'][i]))
    for n, (a, b) in zip(bnames, blocks):
        pr('  %-12s 充电 %.4f  放电 %.4f' % (n, r['ch'][a:b].sum(), r['dis'][a:b].sum()))
    pr('  0:00 储电量 %.4f   24:00 储电量 %.4f' % (r['soc'][0], r['soc'][-1]))
    pr('  紧急购电量: %.6f' % r['e'].sum())

# ---------- write result2.xlsx ----------
OUT_START = 31                       # 2025-02-01
out_days = list(range(OUT_START, NDAYS))
shutil.copy('附件/附件5/result2.xlsx', 'result2.xlsx')
wb = load_workbook('result2.xlsx')

# sheet 1: 计划购电量 (rows already exist: 334 dates)
ws = wb['计划购电量']
for i, d in enumerate(out_days):
    r = results[d]
    row = i + 2
    for t in range(T):
        ws.cell(row=row, column=2 + t, value=round(float(r['buy'][t]), 4))
    ws.cell(row=row, column=146, value=round(float(r['buy'].sum()), 4))
    ws.cell(row=row, column=147, value=round(float(r['cost']), 4))

# sheet 2: 充放电量 — rebuild all rows
ws = wb['充放电量']
ws.delete_rows(2, ws.max_row)
row = 2
for d in out_days:
    r = results[d]
    for k, (a, b) in enumerate(blocks):
        if k == 0:
            ws.cell(row=row, column=1, value=dates[d].to_pydatetime())
        ws.cell(row=row, column=2, value=bnames[k])
        ws.cell(row=row, column=3, value=round(float(r['ch'][a:b].sum()), 4))
        ws.cell(row=row, column=4, value=round(float(r['dis'][a:b].sum()), 4))
        if k == 0:
            ws.cell(row=row, column=5, value='0:00')
            ws.cell(row=row, column=6, value=round(float(r['soc'][0]), 4))
        elif k == 1:
            ws.cell(row=row, column=5, value='24:00')
            ws.cell(row=row, column=6, value=round(float(r['soc'][-1]), 4))
        row += 1

# sheet 3: 紧急购电量 — only days/periods with nonzero emergency
ws = wb['紧急购电量']
ws.delete_rows(2, ws.max_row)
row = 2
n_emerg_rows = 0
for d in out_days:
    r = results[d]
    idx = np.where(r['e'] > 1e-6)[0]
    if len(idx) == 0:
        continue
    first = True
    for t in idx:
        if first:
            ws.cell(row=row, column=1, value=dates[d].to_pydatetime())
            first = False
        h, m = t // 6, (t % 6) * 10
        h2, m2 = (t + 1) // 6, ((t + 1) % 6) * 10
        lab = '%d:%02d-%d:%02d' % (h, m, h2, m2)
        ws.cell(row=row, column=2, value=lab)
        ws.cell(row=row, column=3, value=round(float(r['e'][t]), 4))
        row += 1
        n_emerg_rows += 1
pr('\n紧急购电量 sheet 写入行数: %d' % n_emerg_rows)

wb.save('result2.xlsx')
pr('result2.xlsx 已生成')
log.close()
